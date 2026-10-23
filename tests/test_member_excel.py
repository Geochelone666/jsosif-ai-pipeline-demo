"""Reviewed paper responses can be imported from Excel without activating email."""
from copy import deepcopy
from datetime import time, datetime
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook, load_workbook
import pytest

from create_member_template import create
from subscriptions import CSV_FIELDS, import_members_excel


def test_excel_template_is_importable_disabled_and_retains_survey_preferences(tmp_path):
    path = tmp_path / 'members.xlsx'
    create(path)
    config = import_members_excel({'groups': {}, 'members': []}, path)
    member = config['members'][0]
    assert not member['enabled']
    assert member['send_time'] == '08:30'
    assert member['send_weekday'] == 2
    assert member['preferred_frequency'] == 'Weekly'
    assert member['preferred_time'] == 'Before 9 AM'
    assert member['content_preferences'] == 'Metrics;News'
    assert member['feedback'].startswith('Example only.')


def test_excel_survey_answers_do_not_override_the_reviewed_weekly_schedule(tmp_path):
    path = tmp_path / 'members.xlsx'
    create(path)
    workbook = load_workbook(path)
    sheet = workbook['Members']
    sheet['J2'] = 'Daily'
    sheet['K2'] = 'Monday;Wednesday;Friday'
    sheet['L2'] = 'After 5 PM'
    sheet['H2'] = 4
    sheet['I2'] = time(17, 30)
    workbook.save(path)
    config = import_members_excel({'groups': {}, 'members': []}, path)
    member = config['members'][0]
    assert member['preferred_frequency'] == 'Daily'
    assert member['preferred_weekdays'] == 'Monday;Wednesday;Friday'
    assert member['send_weekday'] == 4 and member['send_time'] == '17:30'
    assert not member['enabled']


def test_excel_formulas_cannot_be_imported_or_replace_existing_configuration(tmp_path):
    path = tmp_path / 'members.xlsx'
    create(path)
    workbook = load_workbook(path)
    workbook['Members']['B2'] = '=CONCAT("Example", " Analyst")'
    workbook.save(path)
    config = {'groups': {}, 'members': []}
    previous = deepcopy(config)
    with pytest.raises(ValueError, match='formulas'):
        import_members_excel(config, path)
    assert config == previous


def test_excel_import_handles_a_single_unnamed_sheet_and_blank_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(CSV_FIELDS))
    sheet.append(['one', 'Example Analyst', 'analyst@example.com', '', 'CSCO',
                  True, 'America/Toronto', 2, time(9, 15)])
    sheet.append([None] * len(CSV_FIELDS))
    path = tmp_path / 'members.xlsx'
    workbook.save(path)
    config = import_members_excel({'groups': {}, 'members': []}, path)
    assert len(config['members']) == 1
    assert config['members'][0]['send_time'] == '09:15'


def test_excel_cli_creates_owner_only_json(tmp_path):
    path = tmp_path / 'members.local.xlsx'
    create(path)
    target = tmp_path / 'subscriptions.local.json'
    root = Path(__file__).resolve().parents[1]
    subprocess.run([sys.executable, str(root / 'scripts/subscriptions.py'), str(target),
                    '--import-excel', str(path)], capture_output=True, check=True)
    assert target.stat().st_mode & 0o777 == 0o600
    assert not json.loads(target.read_text())['members'][0]['enabled']


def test_excel_start_date_preserves_the_future_delivery_gate(tmp_path):
    from deliver_weekly import due_period
    from datetime import timezone

    path = tmp_path / 'members.xlsx'
    create(path)
    workbook = load_workbook(path)
    sheet = workbook['Members']
    sheet['H2'] = 0
    sheet['I2'] = time(9)
    sheet['O2'] = datetime(2026, 10, 26)
    workbook.save(path)
    member = import_members_excel({'groups': {}, 'members': []}, path)['members'][0]
    assert member['delivery_start_date'] == '2026-10-26'
    assert due_period(member, datetime(2026, 10, 23, 14, tzinfo=timezone.utc)) is None
    assert due_period(member, datetime(2026, 10, 26, 13, tzinfo=timezone.utc)) == '2026-10-26'


def test_excel_start_date_rejects_time_components(tmp_path):
    path = tmp_path / 'members.xlsx'
    create(path)
    workbook = load_workbook(path)
    workbook['Members']['O2'] = datetime(2026, 10, 26, 12)
    workbook.save(path)
    with pytest.raises(ValueError, match='must not include a time'):
        import_members_excel({'groups': {}, 'members': []}, path)
