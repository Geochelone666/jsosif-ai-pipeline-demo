"""Private member/group configuration, including a CSV import for manual entry."""
import argparse
import csv
from copy import deepcopy
import io
from datetime import time, date
import json
import re
import os
from tempfile import NamedTemporaryFile
from pathlib import Path
from zoneinfo import ZoneInfo

from email_delivery import validate_recipients
from universe import SUPPORTED


def validate(config):
    if not isinstance(config, dict) or not isinstance(config.get('groups'), dict):
        raise ValueError('Expected groups object and members list')
    groups, members = config['groups'], config.get('members')
    if not isinstance(members, list):
        raise ValueError('members must be a list')
    for key, group in groups.items():
        if not isinstance(key, str) or not key.strip() or not isinstance(group, dict):
            raise ValueError('Each group needs a nonblank id and an object')
        if 'name' in group and not isinstance(group['name'], str):
            raise ValueError('Group names must be strings')
        tickers = group.get('tickers')
        if not isinstance(tickers, list) or not tickers or any(not isinstance(t, str) or t not in SUPPORTED for t in tickers):
            raise ValueError(f'Group {key}: tickers must be supported symbols')
    ids, emails = set(), set()
    for member in members:
        if not isinstance(member, dict):
            raise ValueError('Each member must be an object')
        for field in ('id', 'name', 'email', 'timezone', 'send_time'):
            value = member.get(field)
            if not isinstance(value, str) or not value.strip() or any(ord(c) < 32 for c in value):
                raise ValueError(f'Member {field} must be nonblank text without control characters')
        validate_recipients([member['email']])
        if member['id'] in ids or member['email'].lower() in emails:
            raise ValueError('Duplicate member id or email; combine group memberships into one record')
        ids.add(member['id']); emails.add(member['email'].lower())
        if type(member.get('enabled')) is not bool:
            raise ValueError('enabled must be true or false')
        if type(member.get('send_weekday')) is not int or not 0 <= member['send_weekday'] <= 6:
            raise ValueError('send_weekday must be 0 (Monday) to 6 (Sunday)')
        if not re.fullmatch(r'(?:[01][0-9]|2[0-3]):[0-5][0-9]', member['send_time']):
            raise ValueError('send_time must be HH:MM')
        try:
            ZoneInfo(member['timezone'])
        except (KeyError, ValueError):
            raise ValueError('timezone must be an IANA time zone') from None
        if 'delivery_start_date' in member:
            value = member['delivery_start_date']
            if not isinstance(value, str) or not re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}', value):
                raise ValueError('delivery_start_date must be YYYY-MM-DD')
            date.fromisoformat(value)
        assigned = member.get('groups')
        requested = member.get('tickers', [])
        if (not isinstance(requested, list) or
                any(not isinstance(t, str) or t not in SUPPORTED for t in requested)):
            raise ValueError('Member tickers must be a list of supported symbols')
        if not isinstance(assigned, list) or any(not isinstance(g, str) or g not in groups for g in assigned):
            raise ValueError('Member groups must be a list of known group ids')
        if not assigned and not requested:
            raise ValueError('Each member needs a group or personal holdings')
    for key in ('allow_degraded',):
        if key in config and type(config[key]) is not bool:
            raise ValueError(f'{key} must be boolean')
    max_age = config.get('max_report_age_days', 4)
    if type(max_age) is not int or not 0 <= max_age <= 30:
        raise ValueError('max_report_age_days must be 0 to 30')
    return config


def load(path):
    return validate(json.loads(Path(path).read_text(encoding='utf-8')))


def tickers_for(config, member=None):
    if member is None:
        return sorted({ticker for m in config['members'] if m['enabled']
                       for ticker in tickers_for(config, m)})
    if member.get('tickers'):
        return sorted(set(member['tickers']))
    return sorted({ticker for group in member['groups'] for ticker in config['groups'][group]['tickers']})


def research_tickers(config):
    """Email opt-out never removes a holding from the shared research website."""
    return sorted({t for group in config['groups'].values() for t in group['tickers']} |
                  {t for member in config['members'] for t in member.get('tickers', [])})


CSV_FIELDS = ('id', 'name', 'email', 'groups', 'tickers', 'enabled', 'timezone',
              'send_weekday', 'send_time', 'preferred_frequency', 'preferred_weekdays', 'preferred_time',
              'content_preferences', 'feedback')
REQUIRED_CSV_FIELDS = {'id', 'name', 'email', 'groups', 'enabled', 'timezone',
                       'send_weekday', 'send_time'}


def import_members_csv(config, text):
    """Validate a complete spreadsheet before replacing any private configuration."""
    if not isinstance(config, dict) or not isinstance(config.get('groups'), dict):
        raise ValueError('CSV import needs a configuration with a groups object')
    result = deepcopy(config)
    reader = csv.DictReader(io.StringIO(text.lstrip('\ufeff')), strict=True)
    headers = reader.fieldnames or []
    required = REQUIRED_CSV_FIELDS  # Older spreadsheets remain supported.
    if (len(headers) != len(set(headers)) or not required <= set(headers) or
            set(headers) - set(CSV_FIELDS)):
        raise ValueError('CSV columns must be: ' + ','.join(CSV_FIELDS))
    try:
        source_rows = list(reader)
    except csv.Error:
        raise ValueError('Malformed CSV quoting') from None
    rows = []
    for number, row in enumerate(source_rows, 2):
        if None in row or any(value is None for value in row.values()):
            raise ValueError(f'CSV row {number}: incorrect number of columns')
        row = {key: value.strip() for key, value in row.items()}
        if row['enabled'].lower() not in ('true', 'false', 'yes', 'no'):
            raise ValueError(f'CSV row {number}: enabled must be true/false or yes/no')
        if not re.fullmatch('[0-6]', row['send_weekday']):
            raise ValueError(f'CSV row {number}: send_weekday must be 0 (Monday) to 6 (Sunday)')
        split = lambda value: list(dict.fromkeys(p.strip() for p in value.split(';') if p.strip()))
        row.update(enabled=row['enabled'].lower() in ('true', 'yes'),
                   send_weekday=int(row['send_weekday']), groups=split(row['groups']),
                   tickers=[t.upper() for t in split(row.get('tickers', ''))])
        rows.append(row)
    # A header-only export must not accidentally erase the current roster.
    if not rows:
        raise ValueError('CSV must contain at least one member')
    for row in rows:
        for group in row['groups']:
            if group not in result['groups']:
                holdings = sorted({t for m in rows if group in m['groups'] for t in m['tickers']})
                if not holdings:
                    raise ValueError('New groups need holdings in the spreadsheet or existing configuration')
                result['groups'][group] = {'name': group, 'tickers': holdings}
    result['members'] = rows
    return validate(result)


def import_members_excel(config, path):
    """Read a reviewed Excel roster; formulas are never evaluated or imported."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook['Members'] if 'Members' in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows()
        headers = [str(cell.value).strip() if cell.value is not None else '' for cell in next(rows, [])]
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(headers)
        for row in rows:
            if all(cell.value is None for cell in row):
                continue
            values = []
            for cell in row:
                if cell.data_type == 'f':
                    raise ValueError('Member spreadsheets must contain values, not formulas')
                value = cell.value
                if isinstance(value, time):
                    if value.second or value.microsecond:
                        raise ValueError('Send times must have minute precision')
                    value = value.strftime('%H:%M')
                values.append('' if value is None else str(value))
            writer.writerow(values)
        return import_members_csv(config, stream.getvalue())
    finally:
        workbook.close()


def write_private_config(path, config):
    """Atomically replace a private roster with owner-only permissions."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with NamedTemporaryFile(mode='w', encoding='utf-8', dir=path.parent,
                                prefix='.subscriptions-', delete=False) as stream:
            temporary = Path(stream.name)
            json.dump(validate(config), stream, ensure_ascii=False, indent=2)
            stream.write('\n')
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('config', type=Path)
    parser.add_argument('--import-csv', type=Path, help='Replace members from a manually entered CSV')
    parser.add_argument('--import-excel', type=Path, help='Replace members from a reviewed .xlsx workbook')
    args = parser.parse_args()
    if args.import_csv and args.import_excel:
        parser.error('Choose one member spreadsheet to import')
    if not args.config.exists() and not (args.import_csv or args.import_excel):
        parser.error('Configuration file does not exist; import a reviewed spreadsheet first')
    config = (json.loads(args.config.read_text(encoding='utf-8')) if args.config.exists()
              else {'groups': {}, 'members': []})
    if args.import_csv:
        config = import_members_csv(config, args.import_csv.read_text(encoding='utf-8-sig'))
    if args.import_excel:
        config = import_members_excel(config, args.import_excel)
    validate(config)
    if args.import_csv or args.import_excel:
        write_private_config(args.config, config)
    print(f'Configuration valid: {len(config["members"])} members; {len(tickers_for(config))} active tickers')


if __name__ == '__main__':
    main()
