"""Create the English Excel template for transcribing paper pilot signup slips."""
from datetime import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from subscriptions import CSV_FIELDS


def create(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Members'
    sheet.append(list(CSV_FIELDS))
    sheet.append(['member-001', 'Example Analyst', 'analyst@example.com', 'energy',
                  'CNQ.TO;ENB', False, 'America/Toronto', 2, time(8, 30),
                  'Weekly', 'Wednesday', 'Before 9 AM', 'Metrics;News',
                  'Example only. Replace this row with reviewed member details.'])
    widths = [19, 26, 34, 22, 26, 16, 26, 19, 19, 24, 26, 24, 30, 55]
    for column, width in zip(sheet.columns, widths):
        sheet.column_dimensions[column[0].column_letter].width = width
    for cell in sheet[1]:
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='173B4D' if cell.column <= 9 else '52626E')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet['I2'].number_format = 'hh:mm'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:N{sheet.max_row}'
    enabled = DataValidation(type='list', formula1='"true,false,yes,no"')
    enabled.errorTitle = 'Choose participation'
    enabled.error = 'Use true/yes only after reviewing email participation.'
    enabled.showErrorMessage = True
    enabled.errorStyle = 'stop'
    sheet.add_data_validation(enabled)
    enabled.add('F2:F1000')
    weekday = DataValidation(type='whole', operator='between', formula1=0, formula2=6)
    weekday.showErrorMessage = True
    weekday.error = 'Use Monday = 0 through Sunday = 6.'
    weekday.errorStyle = 'stop'
    sheet.add_data_validation(weekday)
    weekday.add('H2:H1000')

    guide = workbook.create_sheet('Read me')
    guide.append(['Topic', 'Instructions'])
    entries = [
        ('Pilot cadence', 'The pilot sends one email per member per week. Survey frequency preferences do not change the scheduler.'),
        ('Paper slips', 'Transcribe the original answers into the optional preference and feedback columns. Keep operational choices separate.'),
        ('Example row', 'Replace the example before enrollment. It is disabled and uses an example.com address.'),
        ('Stable ID', 'Use a unique stable id for each person. Preserve it when editing names or email addresses.'),
        ('Groups and holdings', 'Use semicolons between groups or ticker symbols. Personal tickers replace group defaults; blank inherits defaults.'),
        ('Participation', 'enabled controls weekly email. A contact-interest checkbox alone does not automatically enable sending.'),
        ('Time zone', 'The paper form uses Toronto time. Use America/Toronto unless the member confirms a different zone.'),
        ('Exact time', 'send_time needs HH:MM. Before 9 AM or After 5 PM belongs in preferred_time until an exact time is confirmed.'),
        ('Weekday', 'send_weekday is one confirmed day: Monday=0, Tuesday=1, Wednesday=2, Thursday=3, Friday=4, Saturday=5, Sunday=6.'),
        ('Optional survey data', 'preferred_frequency, preferred_weekdays, preferred_time, content_preferences and feedback are retained privately, not interpreted as delivery instructions.'),
        ('Privacy', 'Keep the completed workbook private. Do not commit real names or addresses to the repository.'),
        ('Import', 'python scripts/subscriptions.py config/subscriptions.local.json --import-excel config/members.local.xlsx'),
        ('Activation', 'Importing does not send email. Review private previews and enable the workflow separately.'),
    ]
    for row in entries:
        guide.append(row)
    guide.column_dimensions['A'].width = 25
    guide.column_dimensions['B'].width = 110
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for row in range(2, guide.max_row + 1):
        guide.row_dimensions[row].height = 44
    guide.freeze_panes = 'A2'
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


if __name__ == '__main__':
    create(Path(__file__).resolve().parents[1] / 'docs/members.example.xlsx')
